"""The timesheet as a spreadsheet.

The **shape** of this sheet is carried over from Acme's `buildTimeTrackingXlsx`
(`.reference-acme/website/vite.config.js`): a four-row header block with merged cells,
a bold header row on a grey fill, a frozen pane below it, and widths 14 / 10 / 80. None
of that is deducible from a specification -- it is what makes the file usable rather
than merely correct, tuned by years of somebody actually scrolling and filtering it.
The `exceljs` code that produced it is not carried: it was a JavaScript library inside a
Node dev server this product no longer has.

Three things are rewritten:
  * hours are **numbers** with a `number_format`, not pre-formatted strings -- a text
    cell cannot be summed, filtered numerically or charted;
  * the date is a **date**, not a string;
  * the total is a `SUBTOTAL(109, ...)` **formula**, not a constant. Whoever receives a
    timesheet filters it, and a constant total after a filter contradicts the column
    above it -- §6.2's principle ("the printed rows win") applied to a spreadsheet.

Nothing in this module escapes anything. A cell value is not markup, so the description
passes through no escaper at all -- which is the whole point: Acme stored it already
escaped for Typst and wrote that same string into the cell, so clients read
`Call con \\@mario su \\[fase 1\\]` in their spreadsheet.

Pure: primitives in, bytes out, no session and no ORM, so the sheet shape is provable
on its own.
"""

from collections.abc import Sequence
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET_TITLE = "Rapporto ore"
HEADER_FILL_ARGB = "FFD9D9D9"
COLUMN_WIDTHS = (14, 10, 80)
HEADER_ROW = 5
FIRST_DATA_ROW = HEADER_ROW + 1
# `ySplit: 5` in exceljs terms: everything above row 6 stays put while the entries
# scroll.
FREEZE_PANES = f"A{FIRST_DATA_ROW}"
DATE_FORMAT = "dd/mm/yyyy"
HOURS_FORMAT = "0.00"
COLUMN_HEADERS = ("DATA", "ORE", "DESCRIZIONE")


def build_time_report_xlsx(
    *,
    cliente: str,
    offerta: str,
    periodo: str,
    rows: Sequence[tuple[date, Decimal, str]],
) -> bytes:
    workbook = Workbook()
    # `workbook.active` rather than this would be the obvious spelling, but it is typed
    # `... | None` -- correctly so, because a workbook *read from disk* can have no
    # active sheet. A freshly constructed one always has exactly one, and indexing
    # `worksheets` states that fact instead of asserting it away at runtime.
    sheet = workbook.worksheets[0]
    sheet.title = SHEET_TITLE

    bold = Font(bold=True)
    grey = PatternFill(start_color=HEADER_FILL_ARGB, end_color=HEADER_FILL_ARGB, fill_type="solid")

    sheet["A1"] = SHEET_TITLE
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.merge_cells("A1:C1")
    for row, (label, value) in enumerate(
        (("Cliente", cliente), ("Offerta", offerta), ("Periodo", periodo)), start=2
    ):
        sheet[f"A{row}"] = label
        sheet[f"A{row}"].font = bold
        sheet[f"B{row}"] = value
        sheet.merge_cells(f"B{row}:C{row}")

    for index, header in enumerate(COLUMN_HEADERS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=header)
        cell.font = bold
        cell.fill = grey

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for offset, (giorno, ore, descrizione) in enumerate(rows):
        row = FIRST_DATA_ROW + offset
        data_cell = sheet.cell(row=row, column=1, value=giorno)
        data_cell.number_format = DATE_FORMAT
        # `Decimal` handed to openpyxl unchanged: it writes it as a numeric cell, so no
        # float ever exists on this path. Converting to `float` here would reintroduce
        # exactly the drift `Numeric(8,2)` exists to avoid.
        ore_cell = sheet.cell(row=row, column=2, value=ore)
        ore_cell.number_format = HOURS_FORMAT
        # Raw, unescaped, newlines intact. `wrap_text` is what makes a multi-line
        # description readable in the 80-wide column instead of a single clipped line.
        descrizione_cell = sheet.cell(row=row, column=3, value=descrizione)
        descrizione_cell.alignment = Alignment(wrap_text=True, vertical="top")

    last_data_row = HEADER_ROW + len(rows)
    total_row = last_data_row + 1
    label_cell = sheet.cell(row=total_row, column=1, value="Totale ore")
    label_cell.font = bold
    # 109 is SUM-ignoring-hidden-rows: the total follows the filter, so it never
    # contradicts the visible column. A constant here is the spreadsheet form of the
    # defect §6.2 rules out.
    total_cell = sheet.cell(
        row=total_row,
        column=2,
        value=f"=SUBTOTAL(109,B{FIRST_DATA_ROW}:B{last_data_row})",
    )
    total_cell.font = bold
    total_cell.number_format = HOURS_FORMAT

    # The header row plus the data rows only -- never the total row, or filtering would
    # hide the very total it is supposed to update.
    sheet.auto_filter.ref = f"A{HEADER_ROW}:C{last_data_row}"
    sheet.freeze_panes = FREEZE_PANES

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
